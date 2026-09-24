from __future__ import annotations

import openpyxl

from src.core.calibration import build_calibration
from src.core.export import export_to_excel
from src.models.project_data import ProjectState
from src.models.series_data import ExtractedPoint, SeriesData
from src.models.types import CombinedMode, ScaleType, SeriesKind


def _series(index: int, kind: SeriesKind, points: list[tuple[float, float]]) -> SeriesData:
    result = SeriesData(index=index, kind=kind)
    result.points = [ExtractedPoint(x=x, y=y) for x, y in points]
    return result


def test_union_combined_does_not_invent_scatter_values(tmp_path) -> None:
    project = ProjectState()
    project.series = [
        _series(1, SeriesKind.DISCRETE, [(0.0, 0.0), (2.0, 2.0)]),
        _series(2, SeriesKind.DISCRETE, [(1.0, 10.0)]),
    ]
    project.combined_mode = CombinedMode.UNION_X
    path = tmp_path / "union.xlsx"
    export_to_excel(project, path)

    workbook = openpyxl.load_workbook(path, data_only=True)
    rows = list(workbook["Combined"].iter_rows(min_row=2, values_only=True))
    assert rows == [(0.0, 0.0, None), (1.0, None, 10.0), (2.0, 2.0, None)]


def test_metadata_uses_actual_calibration_scale(tmp_path) -> None:
    project = ProjectState()
    project.calibration = build_calibration(
        [(0.0, 1.0), (100.0, 100.0)],
        [(100.0, 0.0), (0.0, 1.0)],
        ScaleType.LOG,
        ScaleType.LINEAR,
    )
    # Deliberately stale settings must not mislabel the exported workbook.
    project.settings.x_scale = ScaleType.LINEAR
    path = tmp_path / "metadata.xlsx"
    export_to_excel(project, path)

    workbook = openpyxl.load_workbook(path, data_only=True)
    metadata = dict(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
    assert metadata["x_scale"] == "LOG"
    assert metadata["y_scale"] == "LINEAR"
